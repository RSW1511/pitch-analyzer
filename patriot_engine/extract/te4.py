"""M2 extractor: Technical Exhibit 4 (places of performance and headcount) from a PWS .docx.

Parses the Word table cell by cell (never PDF text; Game Plan rule 3). The declared
TOTAL row is kept separately from the computed column sums so the checks can flag
disagreement (rule 2) and unlocated positions (the RCC Optional Task 2 RCC).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Dict, List, Optional

import docx

COLUMNS = ["pm", "rcc", "rcc_opt2", "tcc", "tsgli_opt1", "rtl_sl", "support"]


@dataclass
class Te4:
    source: str
    sites: List[Dict]
    declared_total: Dict[str, int]
    computed_total: Dict[str, int]
    notes: List[str] = field(default_factory=list)

    def to_dict(self):
        return asdict(self)

    def unlocated(self) -> Dict[str, int]:
        """Positions the TOTAL row counts that no site row carries."""
        return {
            c: self.declared_total.get(c, 0) - self.computed_total.get(c, 0)
            for c in COLUMNS
            if self.declared_total.get(c, 0) != self.computed_total.get(c, 0)
        }


def _row_cells(row) -> List[str]:
    out, prev = [], None
    for c in row.cells:
        if c._tc is prev:
            continue
        prev = c._tc
        out.append(re.sub(r"\s+", " ", c.text).strip())
    return out


def _num(s: str) -> int:
    m = re.search(r"\d+", s or "")
    return int(m.group()) if m else 0


def _classify_header(h: str) -> Optional[str]:
    t = h.lower()
    if "location" in t:
        return "city"
    if "places of performance" in t:
        return "place"
    if "tsgli" in t:
        return "tsgli_opt1"
    if "rcc" in t and "optional" in t:
        return "rcc_opt2"
    if "rcc" in t:
        return "rcc"
    if "tcc" in t:
        return "tcc"
    if "rtl" in t or "sl" in t.split():
        return "rtl_sl"
    if "support" in t:
        return "support"
    if t.strip() == "pm":
        return "pm"
    return None


def extract_te4(path: str | Path) -> Te4:
    d = docx.Document(str(path))
    table = None
    for t in d.tables:
        head = " ".join(_row_cells(t.rows[0])).lower()
        if "places of performance" in head and "rcc" in head and "tcc" in head:
            table = t
            break
    if table is None:
        raise ValueError(f"{path}: no TE4-style staffing table found")

    header = _row_cells(table.rows[0])
    cols = [_classify_header(h) for h in header]
    sites, declared = [], {}
    for r in table.rows[1:]:
        cells = _row_cells(r)
        if not any(cells):
            continue
        rec: Dict = {}
        for name, cell in zip(cols, cells):
            if name is None:
                continue
            rec[name] = cell
        place = rec.get("place", "")
        city = rec.get("city", "")
        counts = {c: _num(rec.get(c, "")) for c in COLUMNS}
        if re.match(r"^\s*total", city, re.I) or re.match(r"^\s*total", place, re.I):
            declared = counts
            continue
        if not place:
            continue
        sites.append({"place": place, "city": city, **counts})

    computed = {c: sum(s[c] for s in sites) for c in COLUMNS}
    notes = []
    for p in d.paragraphs:
        if "no location" in p.text.lower() and "optional task" in p.text.lower():
            notes.append(p.text.strip())
    return Te4(str(path), sites, declared, computed, notes)


def extract_te5_places(path: str | Path) -> List[str]:
    """Places listed in Technical Exhibit 5 (area-of-responsibility table).

    TE4 and TE5 can disagree at release (a site on one list only, which bidders then ask about). The comparison is done by ``questions.py``."""
    d = docx.Document(str(path))
    for t in d.tables:
        head = " ".join(_row_cells(t.rows[0])).lower()
        if "places of performance" in head and "area of responsibility" in head:
            out = []
            for r in t.rows[1:]:
                cells = _row_cells(r)
                if cells and cells[0] and not cells[0].lower().startswith("total"):
                    out.append(cells[0])
            seen, uniq = set(), []
            for p in out:
                if p not in seen:
                    seen.add(p)
                    uniq.append(p)
            return uniq
    return []
