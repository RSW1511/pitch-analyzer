"""M2 extractor: the "Bid Salaries - Wage Rates" workbook (Step 6).

The workbook stacks three tables: SCA positions by site (wage determination wage
and H&W), OCONUS positions priced off a proxy wage, and exempt salaried roles.
Tables are found by their header text rather than fixed cell addresses, so a
re-saved or re-ordered workbook still parses (Game Plan section 11: files get
re-saved by team members).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field, asdict
from decimal import Decimal
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl

from ..pricing.money import D

SCA_HEADER = "sca va functional lcat"
OCONUS_HEADER = "oconus va lcat"
EXEMPT_HEADER = "exempt lcat"


@dataclass
class BidInputs:
    source: str
    sca: List[Dict] = field(default_factory=list)
    oconus: List[Dict] = field(default_factory=list)
    exempt: List[Dict] = field(default_factory=list)

    def to_dict(self):
        return asdict(self)


def _money(v) -> Optional[Decimal]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, Decimal)):
        return D(v)
    s = re.sub(r"[^\d.\-]", "", str(v))
    return D(s) if s else None


def _header_at(ws, r, c0) -> List[Tuple[int, str]]:
    cols = []
    c = c0
    while True:
        v = ws.cell(r, c).value
        if v is None or str(v).strip() == "":
            break
        cols.append((c, re.sub(r"\s+", " ", str(v)).strip()))
        c += 1
    return cols


def _table(ws, r, c0, cols):
    rows = []
    rr = r + 1
    while rr <= ws.max_row:
        first = ws.cell(rr, c0).value
        if first is None or str(first).strip() == "":
            break
        t = str(first).strip().lower()
        if t in (SCA_HEADER, OCONUS_HEADER, EXEMPT_HEADER):
            break
        rows.append({name: ws.cell(rr, c).value for c, name in cols})
        rr += 1
    return rows


def _find(ws, header_text):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip().lower() == header_text:
                return cell.row, cell.column
    return None


def extract_bid_inputs(path: str | Path) -> BidInputs:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    out = BidInputs(source=str(path))
    found = 0
    for ws in wb.worksheets:
        for hdr, bucket in ((SCA_HEADER, "sca"), (OCONUS_HEADER, "oconus"), (EXEMPT_HEADER, "exempt")):
            pos = _find(ws, hdr)
            if not pos:
                continue
            found += 1
            r, c0 = pos
            cols = _header_at(ws, r, c0)
            for raw in _table(ws, r, c0, cols):
                out_row = _normalise(bucket, raw)
                if out_row:
                    getattr(out, bucket).append(out_row)
    if not found:
        raise ValueError(f"{path}: none of the expected bid-salary tables were found")
    return out


def _g(raw: Dict, *keys):
    for k in raw:
        kl = k.lower()
        if any(x in kl for x in keys):
            return raw[k]
    return None


def _normalise(bucket: str, raw: Dict) -> Optional[Dict]:
    lcat = list(raw.values())[0]
    if lcat is None:
        return None
    if bucket == "sca":
        return {
            "labor_category": str(lcat).strip(),
            "scls_title": str(_g(raw, "scls lcat") or "").strip(),
            "site": str(_g(raw, "installation") or "").strip(),
            "city_state": str(_g(raw, "city") or "").strip(),
            "county": str(_g(raw, "county") or "").strip(),
            "wd": str(_g(raw, "wage determination") or "").strip(),
            "wd_rev": str(_g(raw, "revision") or "").strip(),
            "wage": str(_money(_g(raw, "wage rate"))),
            "holidays": int(_money(_g(raw, "holidays")) or 0),
            "hw": str(_money(_g(raw, "h\\&w", "h&w"))),
        }
    if bucket == "oconus":
        return {
            "labor_category": str(lcat).strip(),
            "site": str(_g(raw, "location") or "").strip(),
            "fte": int(_money(_g(raw, "fte")) or 1),
            "salary": str(_money(_g(raw, "salary"))),
            "wage": str(_money(_g(raw, "wage rate"))),
        }
    return {
        "labor_category": str(lcat).strip(),
        "site": str(_g(raw, "location") or "").strip(),
        "fte": int(_money(_g(raw, "fte")) or 1),
        "salary": str(_money(_g(raw, "salary"))),
    }
