"""M2 extractor: read a *filled* government price matrix (Vol II) into plain rows.

Used (a) to calibrate unknown inputs from a prior bid, (b) as ground truth for the
"compare to truth" phase, and (c) to check what the exporter wrote. Reads cached
cell values, so it is only as current as the last time Excel/LibreOffice saved it.
"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Dict, List, Optional

import openpyxl
from openpyxl.utils import column_index_from_string as ci

from ..pricing.money import D


def _val(ws, col: str, row: int):
    return ws.cell(row, ci(col)).value


def _dec(v) -> Optional[Decimal]:
    if v is None or v == "":
        return None
    try:
        return D(v)
    except Exception:
        return None


def _labor_rows(ws, m, label="Labor Sub-Total"):
    cols = m["columns"]
    rows = []
    r = m["first_row"]
    while r <= ws.max_row:
        a = _val(ws, "A", r)
        if isinstance(a, str) and a.strip().startswith(label):
            return rows, r
        cat = _val(ws, cols["category"], r)
        if cat:
            rows.append((r, cat))
        r += 1
    return rows, None


def read_matrix(path: str | Path, mapping: dict) -> Dict:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    out: Dict = {"labor": {}, "optional": {}, "scls": [], "summary": {}}

    lm = mapping["labor"]
    c = lm["columns"]
    for period, sheet in lm["sheets"].items():
        ws = wb[sheet]
        rows, total_row = _labor_rows(ws, lm)
        recs = []
        for r, cat in rows:
            recs.append(
                {
                    "row": r,
                    "category": cat,
                    "place": _val(ws, c["place"], r),
                    "pws": _val(ws, c["pws"], r),
                    "scls": bool(_val(ws, c["scls"], r)),
                    "fte": _dec(_val(ws, c["fte"], r)),
                    "hours": _dec(_val(ws, c["hours"], r)),
                    "rate": _dec(_val(ws, c["rate"], r)),
                    "ext": _dec(_val(ws, c["ext"], r)),
                }
            )
        out["labor"][period] = {
            "rows": recs,
            "total": _dec(_val(ws, c["ext"], total_row)) if total_row else None,
            "total_row": total_row,
        }

    om = mapping["optional"]
    for task, sheet in om["sheets"].items():
        ws = wb[sheet]
        rows, total_row = _labor_rows(ws, om, label="Total")
        per = {}
        for p, pc in om["period_columns"].items():
            recs = []
            for r, cat in rows:
                recs.append(
                    {
                        "row": r,
                        "category": cat,
                        "place": _val(ws, om["columns"]["place"], r),
                        "scls": bool(_val(ws, om["columns"]["scls"], r)),
                        "fte": _dec(_val(ws, om["columns"]["fte"], r)),
                        "hours": _dec(_val(ws, pc["hours"], r)),
                        "rate": _dec(_val(ws, pc["rate"], r)),
                        "ext": _dec(_val(ws, pc["ext"], r)),
                    }
                )
            per[p] = {"rows": recs, "total": _dec(_val(ws, pc["ext"], total_row)) if total_row else None}
        out["optional"][task] = per

    sm = mapping["scls"]
    sc = sm["columns"]
    ws = wb[sm["sheet"]]
    r = sm["example_row"]
    while r <= ws.max_row:
        cat = _val(ws, sc["category"], r)
        if cat:
            rec = {"row": r}
            for k, col in sc.items():
                v = _val(ws, col, r)
                rec[k] = _dec(v) if k not in ("category", "place", "site", "wd_number", "occupation", "occ_code", "wd_rev") else v
            out["scls"].append(rec)
        r += 1

    ws = wb[mapping["summary"]["sheet"]]
    for period, cells in mapping["summary"]["clin_cells"].items():
        out["summary"][period] = {k: _dec(ws[v].value) for k, v in cells.items()}
    out["summary"]["grand_total"] = _dec(ws[mapping["summary"]["grand_total"]].value)
    return out
