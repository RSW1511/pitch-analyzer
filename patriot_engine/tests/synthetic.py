"""Synthetic fixtures. No company or government document is committed to git; these
reproduce only the *structure* the engine depends on, with invented numbers."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import docx
import openpyxl
from openpyxl.styles import Font

from patriot_engine.build_case import build_case
from patriot_engine.extract.bid_inputs import BidInputs
from patriot_engine.extract.te4 import Te4
from patriot_engine.pricing.model import Scenario

CFG = {
    "name": "TOY",
    "solicitation": "TOY-001",
    "periods": [
        {"key": "base", "months": 10, "label": "base"},
        {"key": "oy1", "months": 12},
        {"key": "oy2", "months": 12},
    ],
    "odc": {"base": 1000, "oy1": 1000, "oy2": 1000},
    "odc_optional": {"base": 100, "oy1": 100, "oy2": 100},
    "extension_months": 6,
    "extension_includes_odc": False,
    "oconus_places": ["tokyo"],
    "site_aliases": {},
    "roles": {
        "pm": {"labor_category": "Program Manager", "pws_ref": "5.1", "lcat_id": 1, "salary_match": "Program Manager"},
        "rtl": {"labor_category": "Team Lead", "pws_ref": "5.2", "lcat_id": 2, "salary_match": "Team Lead"},
        "support_lead": {"labor_category": "Team Lead", "pws_ref": "5.3", "lcat_id": 2, "salary_match": "Support Lead"},
        "support": {"labor_category": "Support Coordinator", "pws_ref": "5.4", "lcat_id": 3, "salary_match": "Support Coordinator"},
        "rcc": {"labor_category": "Advocate", "pws_ref": "5", "lcat_id": 4},
        "tcc": {"labor_category": "Advocate", "pws_ref": "5", "lcat_id": 4},
        "rcc_opt2": {"labor_category": "Advocate", "pws_ref": "5.11.2", "lcat_id": 4},
        "tsgli_adj": {"labor_category": "Adjudicator", "pws_ref": "5.11.1", "lcat_id": 5},
        "tsgli_cns": {"labor_category": "Advocate", "pws_ref": "5.11.1", "lcat_id": 4},
    },
    "splits": {"rtl_sl": {"rtl": 1, "support_lead": 1}, "tsgli_opt1": {"tsgli_adj": 1, "tsgli_cns": 1}},
}

LEVERS = {
    "rate.overhead": "0.02", "rate.ga": "0.03", "rate.fee": "0.06", "rate.fringe": "0.30",
    "hours.sca_total": "2080", "hours.sca_vacation": "80", "hours.sca_sick": "56",
    "hours.exempt_billed": "1912", "hours.exempt_divisor": "2080",
    "exempt.adder": "5000", "esc.exempt.oy1": "0.02", "esc.exempt.oy2": "0.025",
}


def toy_te4(sites=None) -> Te4:
    sites = sites or [
        {"place": "Alpha AFB", "city": "Alpha, TX", "pm": 1, "rcc": 1, "rcc_opt2": 0, "tcc": 2, "tsgli_opt1": 2, "rtl_sl": 2, "support": 2},
        {"place": "Bravo AFB", "city": "Bravo, NM", "pm": 0, "rcc": 1, "rcc_opt2": 0, "tcc": 0, "tsgli_opt1": 0, "rtl_sl": 0, "support": 0},
        {"place": "Charlie AFB", "city": "Charlie, FL", "pm": 0, "rcc": 1, "rcc_opt2": 0, "tcc": 0, "tsgli_opt1": 0, "rtl_sl": 0, "support": 0},
        {"place": "Tokyo AB", "city": "Tokyo, JA", "pm": 0, "rcc": 1, "rcc_opt2": 0, "tcc": 0, "tsgli_opt1": 0, "rtl_sl": 0, "support": 0},
    ]
    cols = ["pm", "rcc", "rcc_opt2", "tcc", "tsgli_opt1", "rtl_sl", "support"]
    comp = {c: sum(s[c] for s in sites) for c in cols}
    declared = {**comp, "rcc_opt2": 1}
    return Te4("synthetic", sites, declared, comp)


def toy_inputs() -> BidInputs:
    sca = []
    for site, wage in (("Alpha AFB", "40.00"), ("Bravo AFB", "42.00"), ("Charlie AFB", "38.50")):
        sca.append({"labor_category": "Recovery Care Coordinator (RCC)", "scls_title": "Advocate", "site": site, "city_state": "", "county": "", "wd": "2014-0042", "wd_rev": "39", "wage": wage, "holidays": 11, "hw": "4.93"})
    return BidInputs(
        "synthetic",
        sca=sca,
        oconus=[{"labor_category": "RCC", "site": "Tokyo AB", "fte": 1, "salary": "90000", "wage": "43.27"}],
        exempt=[
            {"labor_category": "Program Manager", "site": "Alpha AFB", "fte": 1, "salary": "100000"},
            {"labor_category": "Team Lead", "site": "Alpha AFB", "fte": 1, "salary": "95000"},
            {"labor_category": "Support Lead", "site": "Alpha AFB", "fte": 1, "salary": "95000"},
            {"labor_category": "Support Coordinator", "site": "Alpha AFB", "fte": 2, "salary": "60000"},
        ],
    )


def toy_case(burden="0.09"):
    te4 = toy_te4()
    case, rep = build_case(
        CFG, te4, toy_inputs(), salaries={"tsgli_adj": Decimal("110000"), "tsgli_cns": Decimal("90000")},
        state_burden={s: Decimal(burden) for s in ("TX", "NM", "FL", "UNK")},
    )
    return case, rep, te4


def toy_scenario(**over) -> Scenario:
    v = dict(LEVERS)
    v.update(over)
    return Scenario("Base", v)


# --- a synthetic government price-matrix template with the same layout -----------------
def make_template(path: str | Path) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Instruction for Offerors")["A1"] = "Instructions"
    s = wb.create_sheet("SUMMARY")
    s["A1"] = "='Instruction for Offerors'!A1"
    for r, v in ((4, "Offeror Name:"), (5, "Cage Code:"), (6, "Version:"), (7, "Date:")):
        s[f"A{r}"] = v
    for rows, name in (((13, 15, 17, 19, 25, 27), "base"), ((31, 33, 35, 37, 39, 41), "oy1"), ((45, 47, 49, 51, 53, 55), "oy2")):
        s[f"K{rows[3]}"] = 1000
        s[f"K{rows[4]}"] = 100
    s["A11"] = "CLIN 0001"
    s["A21"] = "CLIN 0003"
    s["A23"] = "CLIN 0004"
    s["A25"] = "CLIN 0005"
    s["A15"] = "CLIN 0003"
    s["A17"] = "CLIN 0004"
    s["K27"] = "=SUM(K11:K25)"
    s["K41"] = "=SUM(K31:K39)"
    s["K55"] = "=SUM(K45:K53)"
    s["K58"] = "=+K55+K41+K27"
    for sheet in ("Labor Base Transition In", "Labor Option Period 1", "Labor Base Full Performance", "Labor Option Period 2"):
        ws = wb.create_sheet(sheet)
        ws["A15"], ws["B15"], ws["H15"] = "Labor Category ID #", "Labor Category", "Place of Performance"
        ws["A16"], ws["B16"], ws["H16"], ws["N16"], ws["P16"] = "ABC123", "Administrative Assistant", "Fort Example", 520, 45.48
        ws["Q16"] = "=N16*P16"
        for r in range(17, 31):
            ws[f"O{r}"] = "Hr."
            ws[f"Q{r}"] = f"=+N{r}*P{r}"
            ws.merge_cells(f"B{r}:G{r}")
        ws["A31"], ws["N31"], ws["Q31"] = "Labor Sub-Total", "=SUM(N17:N30)", "=SUM(Q17:Q30)"
        ws.merge_cells("B15:G15")
        ws["A32"] = "note below totals"
        ws.merge_cells("A33:C33")
    for sheet in ("Optional Task 1 (PWS 5.11.1)", "Optional Task 2 (PWS 5.11.2)"):
        ws = wb.create_sheet(sheet)
        ws["B19"], ws["A20"], ws["B20"], ws["N20"], ws["P20"] = "Labor Category", "ABC123", "Admin", 520, 45.07
        for r in range(21, 35):
            for h, rt, ex, u in (("N", "P", "Q", "O"), ("R", "T", "U", "S"), ("V", "X", "Y", "W")):
                ws[f"{u}{r}"] = "Hr."
                ws[f"{ex}{r}"] = f"=+{h}{r}*{rt}{r}"
            ws.merge_cells(f"B{r}:G{r}")
        ws["A35"] = "Total"
        for col in "NQRUVY":
            ws[f"{col}35"] = f"=SUM({col}21:{col}34)"
        ws.merge_cells("A36:L36")
    sc = wb.create_sheet("SCLS Breakout")
    for c, t in zip("ABCDEFGHMNOPQRSTU", ["cat", "place", "site", "wd", "rev", "title", "code", "hours", "wage", "tw", "v", "s", "h", "hwr", "hw", "ind", "rate"]):
        sc[f"{c}11"] = t
    sc["A13"], sc["B13"], sc["H13"], sc["I13"], sc["J13"], sc["K13"], sc["M13"], sc["R13"], sc["T13"] = "Administrative Assistant", "Fort Example", 2080, 80, 56, 96, 22.72, 4.93, 26530.25
    sc["L13"] = "=H13-I13-J13-K13"
    sc["U13"] = "=ROUND((N13+O13+P13+Q13+S13+T13)/L13,2)"
    sc["N13"], sc["O13"], sc["P13"], sc["Q13"], sc["S13"] = "=M13*L13", "=I13*M13", "=J13*M13", "=K13*M13", "=H13*R13"
    pa = wb.create_sheet("Pricing Assumptions")
    pa["A9"], pa["B9"] = "No.", "Explanations"
    for i in range(1, 13):
        pa[f"A{9 + i}"] = i
    wb.properties.creator = "Gov Author"
    wb.properties.lastModifiedBy = "Jane Doe (Student)"
    wb.save(str(path))
    return Path(path)


def make_te4_docx(path: str | Path, sites, total_row=None, extra_section=True, renumber=False, heading_text="Hours of operation are 0730 to 1700."):
    d = docx.Document()
    d.add_heading("PART 1 - GENERAL", level=1)
    d.add_paragraph("1.1 Scope")
    d.add_paragraph("The contractor shall provide services at the places listed.")
    num = "1.3" if renumber else "1.2"
    d.add_paragraph(f"{num} Hours of Operation")
    d.add_paragraph(heading_text)
    hdr = ["PLACES OF PERFORMANCE", "LOCATION (CITY)", "PM", "# RCCs", "# RCCs (Optional Task 2)", "# TCCs", "# TSGLIs (Optional Task 1)", "# RTLs/SL", "SUPPORT"]
    t = d.add_table(rows=1, cols=len(hdr))
    for i, h in enumerate(hdr):
        t.rows[0].cells[i].text = h
    cols = ["pm", "rcc", "rcc_opt2", "tcc", "tsgli_opt1", "rtl_sl", "support"]
    for s in sites:
        row = t.add_row().cells
        row[0].text, row[1].text = s["place"], s["city"]
        for i, c in enumerate(cols):
            row[2 + i].text = str(s.get(c, "") or "")
    tot = total_row or {c: sum(s.get(c, 0) for s in sites) for c in cols}
    row = t.add_row().cells
    row[0].text, row[1].text = "", "TOTAL:"
    for i, c in enumerate(cols):
        row[2 + i].text = str(tot.get(c, ""))
    d.save(str(path))
    return Path(path)
