import shutil
from pathlib import Path

import openpyxl
import pytest
import yaml

from patriot_engine import export as X
from patriot_engine.build_case import build_case
from patriot_engine.extract.bid_inputs import BidInputs
from patriot_engine.pricing.model import price_case
from patriot_engine.pricing.money import D, round_cents

from .synthetic import CFG, make_template, toy_case, toy_inputs, toy_scenario, toy_te4

MAPPING = yaml.safe_load((Path(__file__).parents[1] / "cases" / "example" / "mapping.yaml").read_text())
OFFEROR = {"name": "Toy Co", "cage": "TOY01", "date": "1.1.2026"}


def _big_case(n_sites=20):
    te4 = toy_te4()
    sites = [dict(s) for s in te4.sites]
    inputs = toy_inputs()
    for i in range(n_sites):
        name = f"Site{i:02d} AFB"
        sites.append({"place": name, "city": f"Town{i}, TX", "pm": 0, "rcc": 1, "rcc_opt2": 0, "tcc": 0, "tsgli_opt1": 0, "rtl_sl": 0, "support": 0})
        inputs.sca.append({**inputs.sca[0], "site": name, "wage": str(40 + i * 0.5)})
    from patriot_engine.extract.te4 import Te4

    cols = ["pm", "rcc", "rcc_opt2", "tcc", "tsgli_opt1", "rtl_sl", "support"]
    comp = {c: sum(s[c] for s in sites) for c in cols}
    te4 = Te4("synthetic", sites, {**comp, "rcc_opt2": 1}, comp)
    case, rep = build_case(CFG, te4, inputs, salaries={"tsgli_adj": D(110000), "tsgli_cns": D(90000)}, state_burden={"TX": D("0.09"), "NM": D("0.09"), "FL": D("0.09"), "UNK": D("0.09")})
    return case


@pytest.fixture
def exported(tmp_path):
    tpl = make_template(tmp_path / "template.xlsx")
    case = _big_case(20)  # 11 + 20 = more than the template's 14 labor slots
    s = toy_scenario()
    res = price_case(case, s)
    out = tmp_path / "sub.xlsx"
    rep = X.export_submission(case, res, s, tpl, MAPPING, out, OFFEROR, ["Incumbent: no transition hours."], "abc")
    return tpl, out, case, s, res, rep


def test_rows_grow_past_the_templates_14_slots_and_totals_follow(exported):
    tpl, out, case, s, res, rep = exported
    sheet = MAPPING["labor"]["sheets"]["base"]
    assert rep.rows[sheet] > 14
    assert rep.total_rows[sheet] == 17 + rep.rows[sheet]
    wb = openpyxl.load_workbook(out)
    ws = wb[sheet]
    tr = rep.total_rows[sheet]
    assert ws[f"Q{tr}"].value == f"=SUM(Q17:Q{tr - 1})"  # the 14-row SUM trap is fixed
    assert ws[f"M{tr}"].value is not None  # FTEs are totalled (the template never does)


def test_lock_check_passes_and_catches_tampering(exported):
    tpl, out, case, s, res, rep = exported
    assert X.lock_check(tpl, out, MAPPING, rep) == []
    wb = openpyxl.load_workbook(out)
    wb["SUMMARY"]["A9"] = "Tampered"
    wb["Instruction for Offerors"]["A1"] = "Edited instructions"
    wb.save(out)
    diffs = X.lock_check(tpl, out, MAPPING, rep)
    assert any("Instruction for Offerors!A1" in d for d in diffs)


def test_hygiene_strips_links_and_authors(exported):
    tpl, out, *_ = exported
    assert any("Student" in p for p in X.hygiene_check(tpl))  # the blank carries a student's name
    assert X.hygiene_check(out) == []


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_workbook_recalculated_independently_matches_python_to_the_penny(exported):
    tpl, out, case, s, res, rep = exported
    tot = X.recalc_totals(out, MAPPING)
    assert X.penny_match(res, tot) == []
    assert round_cents(tot["grand_total"]) == round_cents(res.matrix_total)


def test_penny_match_reports_mismatch(exported):
    *_, res, rep = exported[0:0] + exported[4:]
    fake = {f"{p}.{k}": D(0) for p in ("base", "oy1", "oy2") for k in ("labor", "opt1", "opt2", "total")}
    fake["grand_total"] = D(0)
    assert X.penny_match(res, fake)


def test_vol3_blanks_every_price_and_keeps_staffing(exported, tmp_path):
    tpl, out, case, s, res, rep = exported
    v3 = tmp_path / "v3.xlsx"
    blanked = X.make_vol3(out, MAPPING, v3, "Toy Co")
    assert blanked
    assert X.vol3_leaks(v3, res) == []
    wb = openpyxl.load_workbook(v3)
    ws = wb[MAPPING["labor"]["sheets"]["base"]]
    assert ws["B17"].value and ws["N17"].value and ws["P17"].value is None and ws["Q17"].value is None
    assert wb["SUMMARY"]["K58"].value is None
    assert X.hygiene_check(v3) == []


def test_template_traps_detected(tmp_path):
    tpl = make_template(tmp_path / "t.xlsx")
    traps = X.template_traps(tpl, MAPPING)
    joined = " ".join(traps)
    assert "duplicate CLIN" in joined and "SUM range fixed at 14" in joined and "never totalled" in joined and "hard-coded" in joined
