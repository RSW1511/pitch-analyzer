"""M0 intake and register.

Any file -> hash, step, version, type. Duplicates are linked, sensitive files are
quarantined, and CUI / source-selection markings switch off AI reading until someone
with authority clears the file (Game Plan 3.5 guardrails: "CUI and access-controlled
attachments stay inside their authorization; check before any AI service sees them").

A pristine, read-only, hash-named copy of every file can be archived at intake
(section 11: several RCC files were re-saved by team members in Excel Online).
"""

from __future__ import annotations

import csv
import hashlib
import os
import re
import shutil
import subprocess
import zipfile
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Dict, Iterable, List, Optional

STEP_RE = re.compile(r"(?:^|[\\/\s_-])(?:step\s*)?([1-8][ABC]?)(?=\s*[-_ .)]|\s+[A-Z])", re.I)
STEP_FOLDER_RE = re.compile(r"step\s*([1-8][ABC]?)\b", re.I)
VERSION_RES = [
    re.compile(r"(\d{1,2}[A-Z]{3}\d{4})"),  # 18AUG2025
    re.compile(r"rev(?:ision)?[ _]?(\d+)", re.I),  # Rev 39
    re.compile(r"amd[_ ]?(\d{4})|amendment[_ ]?(\d{4})", re.I),
    re.compile(r"\bv(\d+(?:\.\d+)?)\b", re.I),
    re.compile(r"(\d{4}[._-]\d{2}[._-]\d{2})"),
]
TYPE_RULES = [
    ("wage_determination", re.compile(r"\bwd[_ ]?\d{4}|wage determination", re.I)),
    ("price_matrix", re.compile(r"price[_ ]matrix|vol(?:ume)?[_ ]?(?:ii|iii)\b|cost[_ ]or[_ ]price", re.I)),
    ("qa", re.compile(r"question[_ ]and[_ ]answer|q\s*&\s*a|_qa|questions", re.I)),
    ("evaluation_plan", re.compile(r"\btoep\b|evaluation plan", re.I)),
    ("pws", re.compile(r"\bpws\b|performance work statement|statement of work|\bsow\b", re.I)),
    ("amendment", re.compile(r"amendment|amd[_ ]?\d{4}", re.I)),
    ("modification", re.compile(r"\bmod(?:ification)?[_ ]?(?:p0)?\d+|\bp000\d+", re.I)),
    ("rfi", re.compile(r"\brfi\b|request for information", re.I)),
    ("solicitation", re.compile(r"rfp|solicitation|\btor\b", re.I)),
    ("protest", re.compile(r"protest|gao|b-4\d{5}", re.I)),
    ("technical_volume", re.compile(r"vol(?:ume)?[_ ]?i\b|technical", re.I)),
    ("debrief", re.compile(r"debrief|successful[_ ]offeror|technical evaluation", re.I)),
    ("salaries", re.compile(r"salar|wage[_ ]rates|compensation", re.I)),
    ("cdrl", re.compile(r"cdrl|exhibit[_ ][a-k]\b", re.I)),
    ("gfp", re.compile(r"\bgfp\b|government furnished", re.I)),
]
SENSITIVE = [
    ("bank details", re.compile(r"routing (number|no)|aba (number|no)|account (number|no)\.?\s*[:#]?\s*\d{6,}|\biban\b|\bswift\b", re.I)),
    ("source selection information", re.compile(r"source selection information|ssi\b[^a-z]|FAR 2\.101 and 3\.104", re.I)),
    ("GAO protective order", re.compile(r"protective order", re.I)),
    ("competitor proprietary", re.compile(r"proprietary(?: and)? confidential|competition sensitive|company proprietary", re.I)),
    ("ssn-like", re.compile(r"\b\d{3}-\d{2}-\d{4}\b")),
]
CUI_RE = re.compile(r"\bCUI\b|controlled unclassified", re.I)
TEXT_EXT = {".txt", ".md", ".csv", ".json", ".html", ".htm", ".eml"}


@dataclass
class Entry:
    path: str
    name: str
    sha256: str
    bytes: int
    step: str = ""
    kind: str = "other"
    version: str = ""
    duplicate_of: str = ""
    cui: bool = False
    quarantine: bool = False
    ai_ok: bool = True
    reasons: List[str] = field(default_factory=list)
    text_read: bool = False

    def to_row(self) -> Dict[str, str]:
        d = asdict(self)
        d["reasons"] = "; ".join(self.reasons)
        return {k: str(v) for k, v in d.items()}


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def infer_step(path: Path, root: Path) -> str:
    rel = str(path.relative_to(root)) if root in path.parents else str(path)
    m = STEP_FOLDER_RE.search(rel)
    if m:
        return m.group(1).upper()
    m = re.search(r"(?:^|[\\/])([1-8][ABC]?)[-_ ]", rel)
    return m.group(1).upper() if m else ""


def infer_version(name: str) -> str:
    for rx in VERSION_RES:
        m = rx.search(name)
        if m:
            return next(g for g in m.groups() if g)
    return ""


def infer_kind(name: str) -> str:
    name = re.sub(r"[_\-.]+", " ", name)  # underscores are word characters to \b
    for k, rx in TYPE_RULES:
        if rx.search(name):
            return k
    return "other"


def extract_text(path: Path, limit: int = 400_000) -> Optional[str]:
    ext = path.suffix.lower()
    try:
        if ext in TEXT_EXT:
            return path.read_text(errors="ignore")[:limit]
        if ext == ".docx":
            import docx

            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for r in t.rows:
                    parts.append(" ".join(c.text for c in r.cells))
            for s in d.sections:
                parts += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
            return "\n".join(parts)[:limit]
        if ext in (".xlsx", ".xlsm"):
            import openpyxl

            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    parts.append(" ".join(str(c) for c in row if c is not None))
                    if i > 2000:
                        break
            return "\n".join(parts)[:limit]
        if ext == ".pdf" and shutil.which("pdftotext"):
            out = subprocess.run(["pdftotext", "-q", "-l", "40", str(path), "-"], capture_output=True, timeout=120)
            return out.stdout.decode("utf8", "ignore")[:limit]
    except Exception:
        return None
    return None


def classify_text(text: str) -> Dict[str, object]:
    reasons, quarantine = [], False
    for label, rx in SENSITIVE:
        if rx.search(text):
            reasons.append(label)
            quarantine = True
    return {"cui": bool(CUI_RE.search(text[:20000])), "quarantine": quarantine, "reasons": reasons}


def intake(root: str | Path, archive_dir: Optional[str | Path] = None, skip_ext: Iterable[str] = (".tmp", ".lnk")) -> List[Entry]:
    root = Path(root)
    entries: List[Entry] = []
    first: Dict[str, str] = {}
    for p in sorted(root.rglob("*")):
        if not p.is_file() or p.suffix.lower() in skip_ext or p.name.startswith("~$"):
            continue
        digest = sha256_file(p)
        e = Entry(
            path=str(p),
            name=p.name,
            sha256=digest,
            bytes=p.stat().st_size,
            step=infer_step(p, root),
            kind=infer_kind(p.name),
            version=infer_version(p.name),
        )
        if digest in first:
            e.duplicate_of = first[digest]
            e.reasons.append("byte-identical duplicate")
        else:
            first[digest] = str(p)
        text = extract_text(p)
        if text:
            e.text_read = True
            c = classify_text(text)
            e.cui, e.quarantine = c["cui"], c["quarantine"]
            e.reasons += [f"contains {r}" for r in c["reasons"]]
        elif p.suffix.lower() == ".pdf":
            e.reasons.append("text not read (pdftotext missing, scanned image, or XFA form): needs OCR or a human")
        if e.cui:
            e.ai_ok = False
            e.reasons.append("CUI marking: no AI service until cleared by compliance")
        if e.quarantine:
            e.ai_ok = False
        entries.append(e)
        if archive_dir:
            dst = Path(archive_dir) / f"{digest[:16]}_{p.name}"
            if not dst.exists():
                dst.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(p, dst)
                os.chmod(dst, 0o444)
    return entries


def write_register(entries: List[Entry], path: str | Path) -> None:
    cols = ["path", "name", "sha256", "bytes", "step", "kind", "version", "duplicate_of", "cui", "quarantine", "ai_ok", "reasons", "text_read"]
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for e in entries:
            w.writerow(e.to_row())


def redact_bank_details(text: str) -> str:
    """Redact account/routing numbers (RCC Mod P00013 Notice of Assignment)."""
    text = re.sub(r"(routing|aba|account)[^\n\d]{0,25}(\d[\d\- ]{5,})", lambda m: f"{m.group(1)} [REDACTED]", text, flags=re.I)
    return text
